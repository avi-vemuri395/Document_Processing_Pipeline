"""
Enhanced document preprocessing with structured extraction and chunking.
Reduces API costs and improves extraction accuracy.
"""

import pdfplumber
import pandas as pd
import numpy as np
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
from dataclasses import dataclass
import re
import json
from PIL import Image
import io

try:
    import fitz  # PyMuPDF for better PDF handling
    PYMUPDF_AVAILABLE = True
except ImportError:
    PYMUPDF_AVAILABLE = False


@dataclass
class PreprocessedChunk:
    """Represents a preprocessed document chunk."""
    chunk_id: str
    chunk_type: str  # 'table', 'text', 'form', 'header', 'footer'
    content: Any  # Could be text, DataFrame, or structured data
    page_number: int
    confidence: float
    metadata: Dict[str, Any]
    requires_ocr: bool = False


@dataclass
class PreprocessedDocument:
    """Complete preprocessed document with all chunks."""
    document_path: Path
    document_type: str
    chunks: List[PreprocessedChunk]
    structured_data: Dict[str, Any]  # Extracted structured elements
    tables: List[pd.DataFrame]
    key_value_pairs: Dict[str, str]
    total_pages: int
    preprocessing_time: float
    needs_llm: bool  # Whether LLM processing is needed


class DocumentPreprocessor:
    """
    Advanced document preprocessing to extract structure before LLM processing.
    Reduces costs and improves accuracy by pre-extracting structured data.
    """
    
    def __init__(self):
        """Initialize preprocessor with patterns and configurations."""
        self.table_confidence_threshold = 0.7
        self.text_chunk_size = 1000  # Characters per chunk
        self.max_chunks_per_document = 20
        
        # Financial patterns for quick extraction
        self.financial_patterns = {
            'currency': re.compile(r'\$[\d,]+\.?\d*'),
            'percentage': re.compile(r'\d+\.?\d*\s*%'),
            'date': re.compile(r'\d{1,2}[/-]\d{1,2}[/-]\d{2,4}'),
            'ssn': re.compile(r'\d{3}-\d{2}-\d{4}'),
            'ein': re.compile(r'\d{2}-\d{7}'),
            'phone': re.compile(r'[\(]?\d{3}[\)]?[-.\s]?\d{3}[-.\s]?\d{4}'),
            'email': re.compile(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}')
        }
        
        # Key field patterns for PFS and other forms
        self.field_patterns = {
            'total_assets': [r'total\s+assets', r'assets\s+total'],
            'total_liabilities': [r'total\s+liabilities', r'liabilities\s+total'],
            'net_worth': [r'net\s+worth', r'total\s+net\s+worth'],
            'first_name': [r'first\s+name', r'given\s+name'],
            'last_name': [r'last\s+name', r'surname', r'family\s+name'],
            'business_name': [r'business\s+name', r'company\s+name', r'dba'],
        }
    
    def preprocess_document(self, file_path: Path) -> PreprocessedDocument:
        """
        Main preprocessing entry point.
        
        Args:
            file_path: Path to document
            
        Returns:
            PreprocessedDocument with extracted structure
        """
        import time
        start_time = time.time()
        
        if file_path.suffix.lower() == '.pdf':
            result = self._preprocess_pdf(file_path)
        elif file_path.suffix.lower() in ['.xlsx', '.xls']:
            result = self._preprocess_excel(file_path)
        else:
            result = self._create_minimal_preprocessing(file_path)
        
        result.preprocessing_time = time.time() - start_time
        return result
    
    def _preprocess_pdf(self, file_path: Path) -> PreprocessedDocument:
        """
        Advanced PDF preprocessing with structure extraction.
        """
        chunks = []
        tables = []
        key_value_pairs = {}
        structured_data = {}
        
        with pdfplumber.open(file_path) as pdf:
            total_pages = len(pdf.pages)
            
            for page_num, page in enumerate(pdf.pages[:10], 1):  # Process first 10 pages
                # Extract text
                text = page.extract_text() or ""
                
                # Extract tables
                page_tables = page.extract_tables()
                if page_tables:
                    for table_idx, table in enumerate(page_tables):
                        if self._is_valid_table(table):
                            df = self._table_to_dataframe(table)
                            tables.append(df)
                            
                            # Create table chunk
                            chunks.append(PreprocessedChunk(
                                chunk_id=f"table_{page_num}_{table_idx}",
                                chunk_type="table",
                                content=df,
                                page_number=page_num,
                                confidence=self._calculate_table_confidence(df),
                                metadata={'rows': len(df), 'cols': len(df.columns)}
                            ))
                            
                            # Extract key-value pairs from table
                            kv_pairs = self._extract_key_values_from_table(df)
                            key_value_pairs.update(kv_pairs)
                
                # Extract form fields (if any)
                if hasattr(page, 'extract_form_fields'):
                    form_fields = page.extract_form_fields()
                    if form_fields:
                        chunks.append(PreprocessedChunk(
                            chunk_id=f"form_{page_num}",
                            chunk_type="form",
                            content=form_fields,
                            page_number=page_num,
                            confidence=0.9,
                            metadata={'field_count': len(form_fields)}
                        ))
                        key_value_pairs.update(form_fields)
                
                # Extract and chunk text
                if text:
                    text_chunks = self._chunk_text(text, page_num)
                    chunks.extend(text_chunks)
                    
                    # Quick extraction of key fields
                    extracted_fields = self._extract_key_fields(text)
                    key_value_pairs.update(extracted_fields)
        
        # Determine if LLM is needed
        needs_llm = self._needs_llm_processing(chunks, key_value_pairs)
        
        # Build structured data
        structured_data = {
            'extracted_fields': key_value_pairs,
            'table_count': len(tables),
            'has_forms': any(c.chunk_type == 'form' for c in chunks),
            'financial_data': self._extract_financial_data(key_value_pairs, tables)
        }
        
        return PreprocessedDocument(
            document_path=file_path,
            document_type='pdf',
            chunks=chunks,
            structured_data=structured_data,
            tables=tables,
            key_value_pairs=key_value_pairs,
            total_pages=total_pages,
            preprocessing_time=0,
            needs_llm=needs_llm
        )
    
    def _preprocess_excel(self, file_path: Path) -> PreprocessedDocument:
        """
        Enhanced Excel preprocessing with structure preservation.
        """
        chunks = []
        tables = []
        key_value_pairs = {}
        
        excel_file = pd.ExcelFile(file_path)
        
        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(file_path, sheet_name=sheet_name)
            
            # Clean the dataframe
            df_cleaned = self._clean_dataframe(df)
            tables.append(df_cleaned)
            
            # Create chunk for sheet
            chunks.append(PreprocessedChunk(
                chunk_id=f"sheet_{sheet_name}",
                chunk_type="table",
                content=df_cleaned,
                page_number=0,
                confidence=0.95,  # High confidence for native Excel
                metadata={
                    'sheet_name': sheet_name,
                    'rows': len(df_cleaned),
                    'cols': len(df_cleaned.columns)
                }
            ))
            
            # Extract key-value pairs
            kv_pairs = self._extract_key_values_from_table(df_cleaned)
            key_value_pairs.update(kv_pairs)
            
            # Look for totals and summaries
            totals = self._extract_totals_from_dataframe(df_cleaned)
            if totals:
                key_value_pairs.update(totals)
        
        structured_data = {
            'sheet_count': len(excel_file.sheet_names),
            'extracted_totals': self._extract_totals_from_dataframe(tables[0]) if tables else {},
            'has_formulas': self._detect_formulas(file_path)
        }
        
        return PreprocessedDocument(
            document_path=file_path,
            document_type='excel',
            chunks=chunks,
            structured_data=structured_data,
            tables=tables,
            key_value_pairs=key_value_pairs,
            total_pages=len(excel_file.sheet_names),
            preprocessing_time=0,
            needs_llm=False  # Excel rarely needs LLM if properly structured
        )
    
    def _chunk_text(self, text: str, page_num: int) -> List[PreprocessedChunk]:
        """
        Intelligently chunk text while preserving context.
        """
        chunks = []
        
        # Split by paragraphs first
        paragraphs = text.split('\n\n')
        
        current_chunk = ""
        chunk_idx = 0
        
        for para in paragraphs:
            if len(current_chunk) + len(para) < self.text_chunk_size:
                current_chunk += para + "\n\n"
            else:
                if current_chunk:
                    chunks.append(PreprocessedChunk(
                        chunk_id=f"text_{page_num}_{chunk_idx}",
                        chunk_type="text",
                        content=current_chunk.strip(),
                        page_number=page_num,
                        confidence=0.8,
                        metadata={'char_count': len(current_chunk)}
                    ))
                    chunk_idx += 1
                current_chunk = para + "\n\n"
        
        # Add remaining chunk
        if current_chunk:
            chunks.append(PreprocessedChunk(
                chunk_id=f"text_{page_num}_{chunk_idx}",
                chunk_type="text",
                content=current_chunk.strip(),
                page_number=page_num,
                confidence=0.8,
                metadata={'char_count': len(current_chunk)}
            ))
        
        return chunks
    
    def _extract_key_fields(self, text: str) -> Dict[str, str]:
        """
        Quick extraction of key fields using patterns.
        """
        extracted = {}
        
        # Extract using patterns
        for field_name, patterns in self.field_patterns.items():
            for pattern in patterns:
                regex = re.compile(f'{pattern}[:\s]*([^\n]+)', re.IGNORECASE)
                match = regex.search(text)
                if match:
                    value = match.group(1).strip()
                    # Clean currency values
                    if '$' in value:
                        value = re.sub(r'[^\d.]', '', value)
                    extracted[field_name] = value
                    break
        
        # Extract special patterns
        ssn_matches = self.financial_patterns['ssn'].findall(text)
        if ssn_matches:
            extracted['ssn'] = ssn_matches[0]
        
        ein_matches = self.financial_patterns['ein'].findall(text)
        if ein_matches:
            extracted['ein'] = ein_matches[0]
        
        return extracted
    
    def _extract_key_values_from_table(self, df: pd.DataFrame) -> Dict[str, str]:
        """
        Extract key-value pairs from tables.
        """
        kv_pairs = {}
        
        # Check if first column contains labels
        if len(df.columns) >= 2:
            for idx, row in df.iterrows():
                if pd.notna(row.iloc[0]) and pd.notna(row.iloc[1]):
                    key = str(row.iloc[0]).strip().lower()
                    value = str(row.iloc[1]).strip()
                    
                    # Map to standard fields
                    for field_name, patterns in self.field_patterns.items():
                        for pattern in patterns:
                            if re.search(pattern, key, re.IGNORECASE):
                                kv_pairs[field_name] = value
                                break
        
        return kv_pairs
    
    def _extract_totals_from_dataframe(self, df: pd.DataFrame) -> Dict[str, Any]:
        """
        Extract total rows from dataframes.
        """
        totals = {}
        
        # Look for rows with 'total' in them
        for idx, row in df.iterrows():
            row_str = ' '.join(str(val) for val in row if pd.notna(val)).lower()
            if 'total' in row_str:
                # Extract numeric values from this row
                for col_idx, val in enumerate(row):
                    if pd.notna(val) and isinstance(val, (int, float)):
                        col_name = df.columns[col_idx] if col_idx < len(df.columns) else f"col_{col_idx}"
                        totals[f"total_{col_name}"] = val
        
        return totals
    
    def _is_valid_table(self, table: List[List]) -> bool:
        """Check if extracted table is valid."""
        if not table or len(table) < 2:
            return False
        if len(table[0]) < 2:
            return False
        # Check if at least 30% cells have content
        total_cells = sum(len(row) for row in table)
        filled_cells = sum(1 for row in table for cell in row if cell)
        return filled_cells / total_cells > 0.3
    
    def _table_to_dataframe(self, table: List[List]) -> pd.DataFrame:
        """Convert table to pandas DataFrame."""
        # Use first row as headers if it looks like headers
        if self._looks_like_header(table[0]):
            df = pd.DataFrame(table[1:], columns=table[0])
        else:
            df = pd.DataFrame(table)
        
        # Clean the dataframe
        return self._clean_dataframe(df)
    
    def _clean_dataframe(self, df: pd.DataFrame) -> pd.DataFrame:
        """Clean and standardize dataframe."""
        # Remove empty rows and columns
        df = df.dropna(how='all')
        df = df.dropna(axis=1, how='all')
        
        # Convert numeric columns
        for col in df.columns:
            # Try to convert to numeric
            try:
                df[col] = pd.to_numeric(df[col].astype(str).str.replace(',', '').str.replace('$', ''), errors='ignore')
            except:
                pass
        
        return df
    
    def _looks_like_header(self, row: List) -> bool:
        """Check if row looks like table headers."""
        if not row:
            return False
        # Headers usually have more text than numbers
        text_count = sum(1 for cell in row if cell and not str(cell).replace('.', '').replace(',', '').isdigit())
        return text_count > len(row) / 2
    
    def _calculate_table_confidence(self, df: pd.DataFrame) -> float:
        """Calculate confidence score for table extraction."""
        if df.empty:
            return 0.0
        
        # Factors: non-null ratio, numeric columns, consistent data types
        non_null_ratio = df.notna().sum().sum() / (len(df) * len(df.columns))
        numeric_cols = sum(1 for dtype in df.dtypes if dtype in ['int64', 'float64'])
        numeric_ratio = numeric_cols / len(df.columns) if len(df.columns) > 0 else 0
        
        confidence = (non_null_ratio * 0.5) + (numeric_ratio * 0.3) + 0.2
        return min(confidence, 1.0)
    
    def _extract_financial_data(self, kv_pairs: Dict, tables: List[pd.DataFrame]) -> Dict[str, Any]:
        """Extract financial data from key-value pairs and tables."""
        financial = {}
        
        # Extract from key-value pairs
        for key, value in kv_pairs.items():
            if any(term in key.lower() for term in ['asset', 'liability', 'income', 'expense', 'worth']):
                try:
                    # Clean and convert to number
                    cleaned = re.sub(r'[^\d.]', '', str(value))
                    if cleaned:
                        financial[key] = float(cleaned)
                except:
                    pass
        
        # Extract from tables
        for df in tables:
            totals = self._extract_totals_from_dataframe(df)
            financial.update(totals)
        
        return financial
    
    def _detect_formulas(self, file_path: Path) -> bool:
        """Detect if Excel file contains formulas."""
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=False)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                            return True
            return False
        except:
            return False
    
    def _needs_llm_processing(self, chunks: List[PreprocessedChunk], kv_pairs: Dict) -> bool:
        """
        Determine if document needs LLM processing.
        """
        # If we extracted most required fields, might not need LLM
        required_fields = ['first_name', 'last_name', 'total_assets', 'total_liabilities']
        extracted_required = sum(1 for field in required_fields if field in kv_pairs)
        
        if extracted_required >= 3:
            return False  # We have enough data
        
        # If tables are incomplete or low confidence
        table_chunks = [c for c in chunks if c.chunk_type == 'table']
        if table_chunks:
            avg_confidence = sum(c.confidence for c in table_chunks) / len(table_chunks)
            if avg_confidence < 0.7:
                return True
        
        # If document is complex (many pages, mixed content)
        if len(chunks) > 10:
            return True
        
        return True  # Default to using LLM
    
    def _create_minimal_preprocessing(self, file_path: Path) -> PreprocessedDocument:
        """Create minimal preprocessing for unsupported formats."""
        return PreprocessedDocument(
            document_path=file_path,
            document_type='unknown',
            chunks=[],
            structured_data={},
            tables=[],
            key_value_pairs={},
            total_pages=1,
            preprocessing_time=0,
            needs_llm=True
        )
    
    def create_optimized_prompt(self, preprocessed: PreprocessedDocument, target_fields: List[str]) -> str:
        """
        Create optimized prompt using preprocessed data.
        """
        prompt = f"Document has been preprocessed. Found {len(preprocessed.key_value_pairs)} fields.\n\n"
        
        if preprocessed.key_value_pairs:
            prompt += "Already extracted fields:\n"
            for key, value in list(preprocessed.key_value_pairs.items())[:10]:
                prompt += f"- {key}: {value}\n"
            prompt += "\n"
        
        if preprocessed.tables:
            prompt += f"Found {len(preprocessed.tables)} tables with financial data.\n\n"
        
        # Only ask for missing fields
        missing_fields = [f for f in target_fields if f not in preprocessed.key_value_pairs]
        if missing_fields:
            prompt += f"Please extract these missing fields: {', '.join(missing_fields)}\n"
        else:
            prompt += "Please verify the extracted values and provide confidence scores.\n"
        
        return prompt